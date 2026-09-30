"""Step 1 - Document extraction (PDF / DOCX -> raw text with source info)."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_MB = 25


class DocumentError(Exception):
    """Raised with a friendly message when a document cannot be used."""


@dataclass
class Page:
    text: str
    page: int | None  # 1-based page number; None for DOCX (no fixed pages)


@dataclass
class LoadedDocument:
    filename: str
    pages: list[Page] = field(default_factory=list)
    page_count: int = 0

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)


def load_pdf(data: bytes, filename: str = "document.pdf") -> LoadedDocument:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("PyMuPDF is not installed. Run: pip install PyMuPDF") from exc

    try:
        pdf = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise DocumentError(f"Could not open '{filename}' as a PDF. The file may be corrupted.") from exc

    pages: list[Page] = []
    try:
        for index, pdf_page in enumerate(pdf, start=1):
            text = pdf_page.get_text("text") or ""
            if text.strip():
                pages.append(Page(text=text, page=index))
        page_count = pdf.page_count
    except Exception as exc:
        raise DocumentError(f"Text extraction failed for '{filename}'.") from exc
    finally:
        pdf.close()

    if not pages:
        raise DocumentError(
            f"No readable text found in '{filename}'. It may be a scanned/image-only PDF "
            "(this workshop app does not include OCR)."
        )
    return LoadedDocument(filename=filename, pages=pages, page_count=page_count)


def load_docx(data: bytes, filename: str = "document.docx") -> LoadedDocument:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("python-docx is not installed. Run: pip install python-docx") from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"Could not open '{filename}' as a DOCX file. It may be corrupted.") from exc

    blocks: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    text = "\n\n".join(blocks)
    if not text.strip():
        raise DocumentError(f"'{filename}' contains no readable text.")
    return LoadedDocument(filename=filename, pages=[Page(text=text, page=None)], page_count=1)


def load_document(data: bytes, filename: str) -> LoadedDocument:
    """Dispatch on file extension. Raises DocumentError with a friendly message."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentError(f"Unsupported file type '{ext or 'unknown'}'. Please upload a PDF or DOCX file.")
    if not data:
        raise DocumentError(f"'{filename}' is empty.")
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise DocumentError(f"'{filename}' is larger than {MAX_FILE_MB} MB. Please use a smaller file.")
    return load_pdf(data, filename) if ext == ".pdf" else load_docx(data, filename)
