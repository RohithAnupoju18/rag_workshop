import io

import pytest

from core.chunker import chunk_document, split_text
from core.document_loader import DocumentError, LoadedDocument, Page, load_document
from core.text_cleaner import clean_text


def _docx_bytes(paragraphs):
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_clean_text_fixes_whitespace_and_line_breaks():
    raw = "Hello   world\nthis is\na broken\n\n\n\nNew paragraph\x00 here-\nwith hyphen"
    out = clean_text(raw)
    assert "\x00" not in out
    assert "Hello world this is a broken" in out
    assert "\n\n\n" not in out
    assert "hyphenated" not in out and "herewith hyphen" in out.replace("here with", "herewith")


def test_clean_text_keeps_bullets_and_punctuation():
    out = clean_text("Requirements:\n- Age 18+\n- Valid ID (required).")
    assert "- Age 18+" in out and "(required)." in out
    assert out.count("\n") == 2


def test_clean_text_empty():
    assert clean_text("") == "" and clean_text("  \n\n ") == ""


def test_chunk_size_and_overlap_respected():
    text = " ".join(f"Sentence number {i} talks about topic {i}." for i in range(80))
    chunks = split_text(text, chunk_size=300, overlap=80)
    assert len(chunks) > 3
    assert all(len(c) <= 300 for c in chunks)
    assert any(chunks[i].split(". ")[-1] in chunks[i + 1] for i in range(len(chunks) - 1))  # overlap exists


def test_invalid_chunk_settings():
    with pytest.raises(ValueError):
        split_text("abc", chunk_size=50, overlap=10)
    with pytest.raises(ValueError):
        split_text("abc " * 100, chunk_size=200, overlap=200)


def test_chunk_metadata():
    doc = LoadedDocument("a.pdf", [Page("First page text. " * 30, 1), Page("Second page text. " * 30, 2)], 2)
    chunks = chunk_document(doc, 300, 50)
    assert {c.page for c in chunks} == {1, 2}
    assert all(c.source == "a.pdf" and c.chunk_id.startswith("a.pdf::p") for c in chunks)
    assert len({c.chunk_id for c in chunks}) == len(chunks)


def test_docx_extraction():
    doc = load_document(_docx_bytes(["Eligibility: age 18 or above.", "Fee: 500 rupees."]), "x.docx")
    assert "Eligibility" in doc.pages[0].text and doc.pages[0].page is None


def test_pdf_extraction():
    fitz = pytest.importorskip("fitz")
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Hello syllabus page one")
    pdf.new_page().insert_text((72, 72), "Second page content")
    doc = load_document(pdf.tobytes(), "s.pdf")
    assert doc.page_count == 2 and doc.pages[1].page == 2 and "Hello" in doc.pages[0].text


def test_image_only_pdf_is_handled_gracefully():
    fitz = pytest.importorskip("fitz")
    pdf = fitz.open()
    pdf.new_page()  # blank page = no text layer
    with pytest.raises(DocumentError):
        load_document(pdf.tobytes(), "scan.pdf")


def test_empty_and_unsupported_files():
    with pytest.raises(DocumentError):
        load_document(b"", "a.pdf")
    with pytest.raises(DocumentError):
        load_document(b"data", "a.txt")
    with pytest.raises(DocumentError):
        load_document(_docx_bytes([]), "empty.docx")
